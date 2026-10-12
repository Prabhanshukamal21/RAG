from io import BytesIO
from pathlib import Path
import fitz
from docx import Document

def load_uploaded_file(upload):
    suffix=Path(upload.name).suffix.lower()
    raw=upload.getvalue()
    if suffix==".pdf":
        with fitz.open(stream=raw,filetype="pdf") as pdf:
            text="\n\n".join(f"[Page {i+1}]\n{p.get_text()}" for i,p in enumerate(pdf))
    elif suffix==".txt":
        text=raw.decode("utf-8",errors="replace")
    elif suffix==".docx":
        doc=Document(BytesIO(raw))
        text="\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
    else: raise ValueError("Supported formats: PDF, TXT, DOCX")
    if not text.strip(): raise ValueError("No text extracted. Scanned PDFs need OCR.")
    return text,{"file_type":suffix[1:]}
